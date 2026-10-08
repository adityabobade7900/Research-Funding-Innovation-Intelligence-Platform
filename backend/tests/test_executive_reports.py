import pytest
import pytest_asyncio
from datetime import datetime, timezone
from httpx import AsyncClient

from tests.conftest import TestingSessionLocal
from app.models.user import User, UserRole
from app.models.profile import Profile
from app.models.publication import Publication, profile_publications
from app.models.patent import Patent, profile_patents
from app.models.funding import FundingOpportunity
from app.services.executive_report_service import ExecutiveReportService


@pytest.mark.asyncio
async def test_executive_dossier_synthesis_structure(client: AsyncClient):
    """
    Validates complete multi-stream dossier synthesis (Innovation, TRL, Commercialization, Grants, Patents).
    """
    async with TestingSessionLocal() as session:
        pub = Publication(
            title="Fault-Tolerant Quantum Error Correction Codes",
            authors="Dr. Qubit, Dr. Algorithm",
            doi="10.1103/prxquantum.2025.101",
            publication_date=datetime(2025, 3, 1, tzinfo=timezone.utc),
            citation_count=45,
            venue="PRX Quantum",
            primary_domain="Quantum Computing",
            source="manual",
        )
        session.add(pub)

        pat = Patent(
            patent_number="US11555666B2",
            title="Surface Code Decoder for Superconducting Quantum Processor",
            assignee="Quantum Scaling Labs",
            inventors="Dr. Qubit",
            filing_date=datetime(2023, 5, 1, tzinfo=timezone.utc),
            publication_date=datetime(2024, 11, 15, tzinfo=timezone.utc),
            patent_classification="G06N 10/70",
            technology_domain="Quantum Computing",
            citation_count=18,
            source="manual",
        )
        session.add(pat)

        funding = FundingOpportunity(
            title="NSF Quantum Leap Challenge Institutes",
            description="Multi-institutional awards for quantum computing scale-up and commercial translation.",
            funding_agency="National Science Foundation",
            funding_amount=5000000.0,
            status="open",
            application_deadline=datetime(2026, 11, 30, tzinfo=timezone.utc),
            external_id="NSF-QLCI-2026",
            source="manual",
        )
        session.add(funding)
        await session.commit()

    resp = await client.get("/api/v1/reports/dossier?domain=Quantum Computing")
    assert resp.status_code == 200
    data = resp.json()["data"]

    assert data["target_name"] == "Quantum Computing"
    assert "report_id" in data
    assert data["report_id"].startswith("DOSSIER-")
    assert "executive_summary" in data
    assert len(data["key_findings"]) >= 3

    # Check synthesis of metrics
    assert data["innovation_score"] > 30.0
    assert data["estimated_trl"] >= 4
    assert data["commercialization_readiness"] > 30.0
    assert "primary_commercial_pathway" in data

    # Check sub-signal metrics
    assert data["publication_metrics"]["total_publications"] >= 1
    assert data["patent_metrics"]["total_patents"] >= 1
    assert data["funding_metrics"]["matching_opportunities_count"] >= 1


@pytest.mark.asyncio
async def test_executive_dossier_strategic_assessment_and_roadmap(client: AsyncClient):
    """
    Validates SWOT strategic assessment quadrants and 3-phase execution roadmap.
    """
    resp = await client.get("/api/v1/reports/dossier?domain=Quantum Computing")
    assert resp.status_code == 200
    data = resp.json()["data"]

    swot = data["strategic_assessment"]
    assert len(swot["strengths"]) >= 1
    assert len(swot["risks_and_bottlenecks"]) >= 1
    assert len(swot["market_opportunities"]) >= 1
    assert len(swot["barriers_to_entry"]) >= 1

    roadmap = data["roadmap"]
    assert len(roadmap) == 3
    phase_names = [p["phase_name"] for p in roadmap]
    assert any("Foundation" in name for name in phase_names)
    assert any("Validation" in name for name in phase_names)
    assert any("Scale" in name for name in phase_names)

    # Validate action item schema in Phase 1
    p1 = roadmap[0]
    assert len(p1["actions"]) >= 1
    assert "owner_role" in p1["actions"][0]
    assert "target_timeline" in p1["actions"][0]
    assert "expected_outcome" in p1["actions"][0]


@pytest.mark.asyncio
async def test_executive_dossier_profile_scoped_isolation(client: AsyncClient):
    """
    Validates that my_profile_only scopes executive dossier strictly to the user's linked assets.
    """
    async with TestingSessionLocal() as session:
        user_exec = User(
            email="exec_researcher@mit.edu",
            hashed_password="hashed_pw_exec",
            full_name="Prof. Sarah Executive",
            role=UserRole.RESEARCHER,
            is_active=True,
        )
        session.add(user_exec)
        await session.flush()

        prof_exec = Profile(user_id=user_exec.id, institution="MIT Media Lab")
        session.add(prof_exec)
        await session.flush()

        pub = Publication(
            title="Neuromorphic Optical Synapses for Edge Computing",
            authors="Prof. Sarah Executive",
            doi="10.1038/s41565-025-001",
            publication_date=datetime(2025, 1, 15, tzinfo=timezone.utc),
            citation_count=80,
            venue="Nature Nanotechnology",
            primary_domain="Neuromorphic Engineering",
            source="manual",
        )
        session.add(pub)
        await session.flush()
        await session.execute(profile_publications.insert().values(profile_id=prof_exec.id, publication_id=pub.id))

        pat = Patent(
            patent_number="US11777888B2",
            title="Phase-Change Optical Memristor Array",
            assignee="MIT",
            inventors="Prof. Sarah Executive",
            filing_date=datetime(2023, 7, 1, tzinfo=timezone.utc),
            publication_date=datetime(2024, 10, 1, tzinfo=timezone.utc),
            patent_classification="G11C 13/00",
            technology_domain="Neuromorphic Engineering",
            citation_count=30,
            source="manual",
        )
        session.add(pat)
        await session.flush()
        await session.execute(profile_patents.insert().values(profile_id=prof_exec.id, patent_id=pat.id))
        await session.commit()

        dossier = await ExecutiveReportService.generate_dossier(profile_id=prof_exec.id, db=session)
        assert dossier.target_type == "PROFILE"
        assert dossier.target_name == "Prof. Sarah Executive"
        assert dossier.innovation_score > 40.0
        assert dossier.publication_metrics["total_publications"] == 1
        assert dossier.patent_metrics["total_patents"] == 1


@pytest.mark.asyncio
async def test_executive_dossier_empty_corpus_and_insufficient_evidence(client: AsyncClient):
    """
    Validates zero-safe handling for domains with no indexed records.
    """
    resp = await client.get("/api/v1/reports/dossier?domain=UnchartedSector999")
    assert resp.status_code == 200
    data = resp.json()["data"]

    assert data["innovation_score"] == 0.0
    assert data["commercialization_readiness"] == 0.0
    assert data["data_sufficiency"] == "INSUFFICIENT_DATA"
    assert data["primary_commercial_pathway"] == "MONITOR_AND_GATHER_EVIDENCE"


@pytest.mark.asyncio
async def test_executive_dossier_export_markdown(client: AsyncClient):
    """
    Validates formatted Markdown dossier export endpoint.
    """
    resp = await client.get("/api/v1/reports/export/markdown?domain=Quantum Computing")
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("text/markdown")
    md_text = resp.text

    assert "# Executive Intelligence Dossier: Quantum Computing" in md_text
    assert "## 1. Executive Summary" in md_text
    assert "## 2. Key Performance Indicators" in md_text
    assert "## 3. Strategic Assessment (SWOT Synthesis)" in md_text
    assert "## 4. Prioritized Commercialization Recommendations" in md_text
    assert "## 5. Strategic Roadmap Timeline" in md_text
    assert "## 6. Governance & Compliance Notice" in md_text


@pytest.mark.asyncio
async def test_executive_dossier_export_json(client: AsyncClient):
    """
    Validates structured JSON dossier export endpoint.
    """
    resp = await client.get("/api/v1/reports/export/json?domain=Quantum Computing")
    assert resp.status_code == 200
    data = resp.json()["data"]

    assert "report_id" in data
    assert "strategic_assessment" in data
    assert "roadmap" in data


@pytest.mark.asyncio
async def test_executive_portfolio_summary_and_benchmarks(client: AsyncClient):
    """
    Validates cross-domain executive summary and portfolio benchmark leaderboard.
    """
    resp = await client.get("/api/v1/reports/summary")
    assert resp.status_code == 200
    data = resp.json()["data"]

    assert "total_domains_benchmarked" in data
    assert "portfolio_average_innovation_score" in data
    assert "portfolio_average_readiness_score" in data
    assert "dominant_pathway" in data
    assert "domain_benchmarks" in data
    assert len(data["domain_benchmarks"]) >= 1


@pytest.mark.asyncio
async def test_executive_reports_api_endpoints(client: AsyncClient):
    """
    Validates all REST endpoints under /api/v1/reports/*.
    """
    # 1. Dossier
    r1 = await client.get("/api/v1/reports/dossier")
    assert r1.status_code == 200

    # 2. Summary
    r2 = await client.get("/api/v1/reports/summary")
    assert r2.status_code == 200

    # 3. Markdown Export
    r3 = await client.get("/api/v1/reports/export/markdown")
    assert r3.status_code == 200

    # 4. JSON Export
    r4 = await client.get("/api/v1/reports/export/json")
    assert r4.status_code == 200


# =========================================================================
# MODULE 11: REPORTS & EXPORT SYSTEM TESTS
# =========================================================================

@pytest.mark.asyncio
async def test_get_report_types(client: AsyncClient):
    """Validates GET /api/v1/reports/types catalog."""
    resp = await client.get("/api/v1/reports/types")
    assert resp.status_code == 200
    data = resp.json()["data"]
    type_ids = [t["id"] for t in data["report_types"]]
    assert "FUNDING" in type_ids
    assert "PATENT" in type_ids
    assert "RESEARCH_TREND" in type_ids
    assert "INNOVATION_INTELLIGENCE" in type_ids
    assert "COMMERCIALIZATION" in type_ids


@pytest.mark.asyncio
async def test_preview_all_report_types(client: AsyncClient, test_user: dict):
    """Validates POST /api/v1/reports/preview across all report categories."""
    from app.core.security import create_access_token
    token = create_access_token({"sub": str(test_user["id"]), "email": test_user["email"]})
    headers = {"Authorization": f"Bearer {token}"}

    # 1. FUNDING Report
    fnd_res = await client.post(
        "/api/v1/reports/preview",
        headers=headers,
        json={"report_type": "FUNDING", "domain": "Quantum Computing"}
    )
    assert fnd_res.status_code == 200
    fnd_data = fnd_res.json()["data"]
    assert fnd_data["report_type"] == "FUNDING"
    assert "Total Grant Pool" in fnd_data["metrics"]

    # 2. PATENT Report
    pat_res = await client.post(
        "/api/v1/reports/preview",
        headers=headers,
        json={"report_type": "PATENT", "domain": "Quantum Computing"}
    )
    assert pat_res.status_code == 200
    pat_data = pat_res.json()["data"]
    assert pat_data["report_type"] == "PATENT"
    assert "Total Tracked Patents" in pat_data["metrics"]

    # 3. RESEARCH_TREND Report
    trend_res = await client.post(
        "/api/v1/reports/preview",
        headers=headers,
        json={"report_type": "RESEARCH_TREND", "domain": "Quantum Computing"}
    )
    assert trend_res.status_code == 200
    trend_data = trend_res.json()["data"]
    assert trend_data["report_type"] == "RESEARCH_TREND"

    # 4. INNOVATION_INTELLIGENCE Report (Must show 5 factors with weights)
    innov_res = await client.post(
        "/api/v1/reports/preview",
        headers=headers,
        json={"report_type": "INNOVATION_INTELLIGENCE", "domain": "Quantum Computing"}
    )
    assert innov_res.status_code == 200
    innov_data = innov_res.json()["data"]
    assert innov_data["report_type"] == "INNOVATION_INTELLIGENCE"
    assert "factors" in innov_data and innov_data["factors"] is not None
    assert len(innov_data["factors"]) == 5
    weights = {f["factor_name"]: f["weight_pct"] for f in innov_data["factors"]}
    assert weights["Research Novelty"] == 30.0
    assert weights["Patent Strength"] == 20.0
    assert weights["Technology Maturity"] == 15.0
    assert weights["Market Potential"] == 20.0
    assert weights["Funding Relevance"] == 15.0

    # 5. COMMERCIALIZATION Report (Conservative language)
    comm_res = await client.post(
        "/api/v1/reports/preview",
        headers=headers,
        json={"report_type": "COMMERCIALIZATION", "domain": "Quantum Computing"}
    )
    assert comm_res.status_code == 200
    comm_data = comm_res.json()["data"]
    assert comm_data["report_type"] == "COMMERCIALIZATION"
    assert "Commercial Readiness Score" in comm_data["metrics"]
    assert "potential" in comm_data["summary_text"].lower()
    assert "guaranteed" not in comm_data["summary_text"].lower()


@pytest.mark.asyncio
async def test_export_pdf_generation_validity(client: AsyncClient, test_user: dict):
    """Validates that POST /api/v1/reports/export/pdf generates a valid, well-formed PDF."""
    from app.core.security import create_access_token
    token = create_access_token({"sub": str(test_user["id"]), "email": test_user["email"]})
    headers = {"Authorization": f"Bearer {token}"}

    resp = await client.post(
        "/api/v1/reports/export/pdf",
        headers=headers,
        json={"report_type": "INNOVATION_INTELLIGENCE", "domain": "Quantum Computing"}
    )
    assert resp.status_code == 200
    assert resp.headers["Content-Type"] == "application/pdf"
    assert "attachment; filename=" in resp.headers["Content-Disposition"]
    content = resp.content
    assert len(content) > 1000
    # PDF magic signature %PDF-
    assert content.startswith(b"%PDF-")


@pytest.mark.asyncio
async def test_export_excel_generation_validity(client: AsyncClient, test_user: dict):
    """Validates that POST /api/v1/reports/export/excel generates a valid openpyxl readable workbook."""
    import io
    import openpyxl
    from app.core.security import create_access_token
    token = create_access_token({"sub": str(test_user["id"]), "email": test_user["email"]})
    headers = {"Authorization": f"Bearer {token}"}

    resp = await client.post(
        "/api/v1/reports/export/excel",
        headers=headers,
        json={"report_type": "INNOVATION_INTELLIGENCE", "domain": "Quantum Computing"}
    )
    assert resp.status_code == 200
    assert "spreadsheetml.sheet" in resp.headers["Content-Type"]
    content = resp.content
    assert len(content) > 1000
    # Standard ZIP / XLSX signature PK
    assert content.startswith(b"PK")

    # Load with openpyxl to verify workbook validity and structure
    wb = openpyxl.load_workbook(io.BytesIO(content))
    assert "Summary & Metrics" in wb.sheetnames
    assert "Data Records" in wb.sheetnames
    assert "5-Pillar Score Factors" in wb.sheetnames
    ws_factors = wb["5-Pillar Score Factors"]
    assert ws_factors.cell(row=2, column=1).value == "Research Novelty"
    assert ws_factors.cell(row=2, column=2).value == 30.0


@pytest.mark.asyncio
async def test_export_unauthenticated_rejected(client: AsyncClient):
    """Validates that export endpoints require authentication."""
    r_pdf = await client.post("/api/v1/reports/export/pdf", json={"report_type": "FUNDING"})
    assert r_pdf.status_code == 401

    r_excel = await client.post("/api/v1/reports/export/excel", json={"report_type": "FUNDING"})
    assert r_excel.status_code == 401


@pytest.mark.asyncio
async def test_export_invalid_report_type(client: AsyncClient, test_user: dict):
    """Validates that an invalid report type is rejected with 422."""
    from app.core.security import create_access_token
    token = create_access_token({"sub": str(test_user["id"]), "email": test_user["email"]})
    headers = {"Authorization": f"Bearer {token}"}

    r_pdf = await client.post(
        "/api/v1/reports/export/pdf",
        headers=headers,
        json={"report_type": "INVALID_TYPE_XYZ"}
    )
    assert r_pdf.status_code == 422

    r_prev = await client.post(
        "/api/v1/reports/preview",
        headers=headers,
        json={"report_type": "NON_EXISTENT"}
    )
    assert r_prev.status_code == 422


@pytest.mark.asyncio
async def test_export_filters_and_empty_results(client: AsyncClient, test_user: dict):
    """Validates filtering with non-matching domain returns clean empty-state preview and exports without crashing."""
    from app.core.security import create_access_token
    token = create_access_token({"sub": str(test_user["id"]), "email": test_user["email"]})
    headers = {"Authorization": f"Bearer {token}"}

    # Filter with non-matching domain
    res = await client.post(
        "/api/v1/reports/preview",
        headers=headers,
        json={
            "report_type": "FUNDING",
            "domain": "NonExistentSpecialtyDomain999",
            "start_date": "2020-01-01",
            "end_date": "2021-01-01",
            "agency": "NASA",
            "min_amount": 50000000.0
        }
    )
    assert res.status_code == 200
    data = res.json()["data"]
    assert data["total_records"] == 0
    assert len(data["table_rows"]) == 0
    assert "No matching records found" in data["summary_text"]

    # Export PDF with empty data must still generate a valid PDF
    pdf_res = await client.post(
        "/api/v1/reports/export/pdf",
        headers=headers,
        json={
            "report_type": "FUNDING",
            "domain": "NonExistentSpecialtyDomain999",
            "min_amount": 50000000.0
        }
    )
    assert pdf_res.status_code == 200
    assert pdf_res.content.startswith(b"%PDF-")

    # Export Excel with empty data must still generate valid workbook
    excel_res = await client.post(
        "/api/v1/reports/export/excel",
        headers=headers,
        json={
            "report_type": "FUNDING",
            "domain": "NonExistentSpecialtyDomain999",
            "min_amount": 50000000.0
        }
    )
    assert excel_res.status_code == 200
    assert excel_res.content.startswith(b"PK")


@pytest.mark.asyncio
async def test_all_five_report_types_export_pdf_and_excel(client: AsyncClient, test_user: dict):
    """Verifies that all 5 required report types generate valid PDF and Excel files."""
    import openpyxl
    import io
    from app.core.security import create_access_token
    token = create_access_token({"sub": str(test_user["id"]), "email": test_user["email"]})
    headers = {"Authorization": f"Bearer {token}"}

    report_types = ["FUNDING", "PATENT", "RESEARCH_TREND", "INNOVATION_INTELLIGENCE", "COMMERCIALIZATION"]
    for rtype in report_types:
        pdf_res = await client.post(
            "/api/v1/reports/export/pdf",
            headers=headers,
            json={"report_type": rtype, "domain": "Artificial Intelligence"}
        )
        assert pdf_res.status_code == 200, f"Failed PDF export for {rtype}"
        assert pdf_res.content.startswith(b"%PDF-"), f"Invalid PDF header for {rtype}"

        excel_res = await client.post(
            "/api/v1/reports/export/excel",
            headers=headers,
            json={"report_type": rtype, "domain": "Artificial Intelligence"}
        )
        assert excel_res.status_code == 200, f"Failed Excel export for {rtype}"
        assert excel_res.content.startswith(b"PK"), f"Invalid Excel ZIP header for {rtype}"
        wb = openpyxl.load_workbook(io.BytesIO(excel_res.content))
        assert "Summary & Metrics" in wb.sheetnames
        assert "Data Records" in wb.sheetnames


